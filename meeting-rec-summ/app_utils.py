"""Local, self-contained replacement for the helpers this tool used to import
from the private ``myfunc`` package (no longer available).

It re-implements, from scratch, exactly the symbols used by ``Openai_Zapisnik.py``
and ``Dunja_Zapisnik.py``:

- from ``myfunc.varvars_dicts``: :func:`work_prompts`, :data:`work_vars`
- from ``myfunc.mojafunkcija``: :func:`positive_login`,
  :func:`initialize_session_state`, :func:`sacuvaj_dokument`
- from ``myfunc.asistenti``: :func:`priprema` and its helpers
  (:func:`transkript`, :func:`read_local_image`, :func:`read_url_image`,
  :func:`generate_corrected_transcript`, :func:`delete_mp3_files`)

The ``asistenti`` helpers preserve the behaviour of the reference copies that
used to live in ``Dunja_Zapisnik.py``. The prompt texts in :func:`work_prompts`
are fresh re-implementations that match how the app uses each key.
"""

from __future__ import annotations

import base64
import glob
import os
from io import BytesIO

import openai
import requests
import streamlit as st
from PIL import Image
from pydub import AudioSegment

try:  # streamlit-authenticator only needed for the Streamlit-cloud login path
    import streamlit_authenticator as stauth
except Exception:  # noqa: BLE001
    stauth = None

try:
    import yaml
    from yaml.loader import SafeLoader
except Exception:  # noqa: BLE001
    yaml = None

CONFIG_PATH = os.path.join(os.path.dirname(__file__), "config.yaml")


# --------------------------------------------------------------------------- #
# myfunc.varvars_dicts                                                         #
# --------------------------------------------------------------------------- #
def work_prompts() -> dict:
    """Return the prompt templates used by the summarizer/transcriber.

    Keys whose values contain ``{...}`` placeholders are consumed by the app via
    ``str.format(...)`` (e.g. ``topic_list_summary`` -> ``number_of_topics``);
    the rest are used as plain system prompts.
    """
    return {
        # Short, overall meeting summary (used for the "Kratak" mode and as intro).
        "intro_summary": (
            "You are an expert meeting-minutes writer. Write a concise, faithful "
            "summary of the meeting transcript provided by the user. Capture the "
            "purpose, the key points discussed, the decisions made and the action "
            "items. Keep it clearly structured."
        ),
        # Date + participants + short intro.
        "date_participants_summary": (
            "From the meeting transcript provided by the user, extract and present: "
            "the meeting date (if mentioned), the list of participants (if "
            "mentioned), and a short 2-3 sentence introduction describing the "
            "purpose of the meeting. If the date or participants are not stated, "
            "say so briefly."
        ),
        # Identify N main topics; the app splits the result on newlines.
        "topic_list_summary": (
            "Read the meeting transcript provided by the user and identify the "
            "{number_of_topics} most important topics discussed. Return ONLY a "
            "plain list with one topic per line, with no numbering, bullets or "
            "extra commentary."
        ),
        # Summary of a single topic.
        "topic_summary": (
            "Summarize what was said in the meeting about the following topic: "
            "{topic}. Base the summary strictly on the transcript provided by the "
            "user; include the relevant details, decisions and action items for "
            "this topic only."
        ),
        # Closing conclusion.
        "conclusion_summary": (
            "Write a short closing conclusion for the meeting based on the "
            "transcript provided by the user: the main outcomes, the agreed next "
            "steps and the overall takeaway."
        ),
        # map_reduce "map" template: must contain {text} and {opis}.
        "summary_begin": (
            "Write a comprehensive summary of the following section of a meeting "
            "transcript.\n\n{opis}\n\nSection:\n{text}\n\nSummary:"
        ),
        # map_reduce "combine" template: must contain {text} and {opis_kraj}.
        "summary_end": (
            "You are given several partial summaries of a long meeting transcript. "
            "Combine them into a single coherent final summary.\n\n{opis_kraj}\n\n"
            "Partial summaries:\n{text}\n\nFinal summary:"
        ),
        # System prompt used when correcting a raw Whisper transcript chunk.
        "text_from_audio": (
            "You are a transcription editor. The user message contains a raw, "
            "automatically generated transcript chunk. Correct obvious spelling, "
            "punctuation and spacing mistakes and fix clearly misrecognized words "
            "WITHOUT changing the meaning and without adding or removing content. "
            "Return only the corrected text."
        ),
        # Default instruction for describing an image.
        "text_from_image": (
            "Describe the contents of this image in detail. If it contains text, "
            "transcribe that text accurately."
        ),
    }


# Subscriptable config dict (the app reads ``work_vars["names"]["openai_model"]``).
work_vars = {
    "names": {
        "openai_model": "gpt-4o-mini",
    }
}

# Module-level prompts, mirroring the original ``mprompts = work_prompts()`` usage.
mprompts = work_prompts()


# --------------------------------------------------------------------------- #
# myfunc.mojafunkcija                                                          #
# --------------------------------------------------------------------------- #
def initialize_session_state(default_values: dict) -> None:
    """Seed ``st.session_state`` with defaults that are not already present."""
    for key, value in default_values.items():
        if key not in st.session_state:
            st.session_state[key] = value


def sacuvaj_dokument(document, file_name: str) -> None:
    """Offer the generated text for download as ``.txt`` and ``.docx``.

    The original helper also produced a PDF via an external ``wkhtmltopdf``
    binary; that system dependency is dropped here to keep the archived tool
    self-contained. Plain-text export always works; DOCX requires ``python-docx``.
    """
    st.download_button(
        "Download .txt",
        data=str(document),
        file_name=f"{file_name}.txt",
        mime="text/plain",
    )
    try:
        from docx import Document

        doc = Document()
        for paragraph in str(document).split("\n"):
            doc.add_paragraph(paragraph)
        buffer = BytesIO()
        doc.save(buffer)
        st.download_button(
            "Download .docx",
            data=buffer.getvalue(),
            file_name=f"{file_name}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    except Exception as exc:  # noqa: BLE001 - DOCX is best-effort
        st.caption(f"DOCX izvoz nije dostupan ({exc}).")


def positive_login(main, name: str = " "):
    """streamlit-authenticator login over ``config.yaml``.

    Runs ``main()`` on success and returns ``(name, authentication_status,
    username)``; mirrors the original ``positive_login(main, " ")`` call shape.
    """
    if stauth is None or yaml is None:
        st.error(
            "streamlit-authenticator/PyYAML nisu instalirani; login nije dostupan."
        )
        st.stop()

    with open(CONFIG_PATH, encoding="utf-8") as handle:
        config = yaml.load(handle, Loader=SafeLoader)

    authenticator = stauth.Authenticate(
        config["credentials"],
        config["cookie"]["name"],
        config["cookie"]["key"],
        config["cookie"]["expiry_days"],
        config.get("preauthorized"),
    )

    user_name, authentication_status, username = authenticator.login("Login", "main")
    if authentication_status:
        with st.sidebar:
            authenticator.logout("Logout", "main", key="logout_btn")
        main()
    elif authentication_status is False:
        st.error("Pogrešno korisničko ime ili lozinka.")
    else:
        st.warning("Unesite korisničko ime i lozinku.")

    return user_name, authentication_status, username


# --------------------------------------------------------------------------- #
# myfunc.asistenti                                                             #
# --------------------------------------------------------------------------- #
def priprema() -> None:
    """Sidebar entry point: pick a preparatory action and run it."""
    izbor_radnji = st.selectbox(
        "Odaberite pripremne radnje",
        (
            "Transkribovanje Zvučnih Zapisa",
            "Čitanje sa slike iz fajla",
            "Čitanje sa slike sa URL-a",
        ),
        help="Odabir pripremnih radnji",
    )
    if izbor_radnji == "Transkribovanje Zvučnih Zapisa":
        transkript()
    elif izbor_radnji == "Čitanje sa slike iz fajla":
        read_local_image()
    elif izbor_radnji == "Čitanje sa slike sa URL-a":
        read_url_image()


def transkript() -> None:
    """Transcribe an uploaded audio/video file and correct the transcript."""
    with st.sidebar:
        st.info("Konvertujte audio/video u TXT")
        audio_file = st.file_uploader(
            "Odaberite audio/video fajl",
            key="audio_",
            help="Odabir dokumenta",
        )
        transcript = ""

        if audio_file is not None:
            st.audio(audio_file.getvalue(), format="audio/mp3")
            placeholder = st.empty()

            with placeholder.form(key="my_jezik", clear_on_submit=False):
                jezik = st.selectbox(
                    "Odaberite jezik izvornog teksta 👉",
                    ("sr", "en"),
                    key="jezik",
                    help="Odabir jezika",
                )

                submit_button = st.form_submit_button(label="Submit")
                client = openai
                if submit_button:
                    with st.spinner("Sačekajte trenutak..."):
                        system_prompt = mprompts["text_from_audio"]
                        transcript = generate_corrected_transcript(
                            client, system_prompt, audio_file, jezik
                        )
                        with st.expander("Transkript"):
                            st.info(transcript)

            if transcript != "":
                st.download_button(
                    "Download transcript",
                    transcript,
                    file_name="transcript.txt",
                    help="Odabir dokumenta",
                )
                delete_mp3_files(".")


def read_local_image() -> None:
    """Describe an image uploaded from a local file."""
    st.info("Čita sa slike")
    image_f = st.file_uploader(
        "Odaberite sliku",
        type="jpg",
        key="slika_",
        help="Odabir dokumenta",
    )
    content = ""

    if image_f is not None:
        base64_image = base64.b64encode(image_f.getvalue()).decode("utf-8")
        image_bytes = base64.b64decode(base64_image)
        image = Image.open(BytesIO(image_bytes))
        st.image(image, width=150)
        placeholder = st.empty()

        with placeholder.form(key="my_image", clear_on_submit=False):
            default_text = mprompts["text_from_image"]
            upit = st.text_area("Unesite uputstvo ", default_text)
            submit_button = st.form_submit_button(label="Submit")

            if submit_button:
                with st.spinner("Sačekajte trenutak..."):
                    api_key = os.getenv("OPENAI_API_KEY")
                    headers = {
                        "Content-Type": "application/json",
                        "Authorization": f"Bearer {api_key}",
                    }
                    payload = {
                        "model": work_vars["names"]["openai_model"],
                        "messages": [
                            {
                                "role": "user",
                                "content": [
                                    {"type": "text", "text": upit},
                                    {
                                        "type": "image_url",
                                        "image_url": {
                                            "url": f"data:image/jpeg;base64,{base64_image}"
                                        },
                                    },
                                ],
                            }
                        ],
                        "max_tokens": 300,
                    }
                    response = requests.post(
                        "https://api.openai.com/v1/chat/completions",
                        headers=headers,
                        json=payload,
                    )
                    json_data = response.json()
                    content = json_data["choices"][0]["message"]["content"]
                    with st.expander("Opis slike"):
                        st.info(content)

        if content != "":
            st.download_button(
                "Download opis slike",
                content,
                file_name=f"{image_f.name}.txt",
                help="Čuvanje dokumenta",
            )


def read_url_image() -> None:
    """Describe an image fetched from a URL."""
    client = openai
    st.info("Čita sa slike sa URL")
    content = ""

    img_url = st.text_input("Unesite URL slike ")
    image_f = os.path.basename(img_url)
    if img_url != "":
        st.image(img_url, width=150)
        placeholder = st.empty()
        with placeholder.form(key="my_image_url", clear_on_submit=False):
            default_text = mprompts["text_from_image"]
            upit = st.text_area("Unesite uputstvo ", default_text)
            submit_button = st.form_submit_button(label="Submit")
            if submit_button:
                with st.spinner("Sačekajte trenutak..."):
                    response = client.chat.completions.create(
                        model=work_vars["names"]["openai_model"],
                        messages=[
                            {
                                "role": "user",
                                "content": [
                                    {"type": "text", "text": upit},
                                    {
                                        "type": "image_url",
                                        "image_url": {"url": img_url},
                                    },
                                ],
                            }
                        ],
                        max_tokens=300,
                    )
                    content = response.choices[0].message.content
                    with st.expander("Opis slike"):
                        st.info(content)

    if content != "":
        st.download_button(
            "Download opis slike",
            content,
            file_name=f"{image_f}.txt",
            help="Čuvanje dokumenta",
        )


def generate_corrected_transcript(client, system_prompt, audio_file, jezik):
    """Convert audio to mono 16 kHz mp3, transcribe in chunks, then GPT-correct."""

    def convert_to_mp3(file_path, output_path):
        audio = AudioSegment.from_file(file_path)
        audio = audio.set_channels(1).set_frame_rate(16000)
        audio.export(output_path, format="mp3", bitrate="128k")

    def transcribe_audio(file_path, jezik):
        with open(file_path, "rb") as audio:
            return client.audio.transcriptions.create(
                model="whisper-1",
                file=audio,
                language=jezik,
                response_format="text",
            )

    def split_mp3_file(
        input_path, output_directory, max_file_size_mb=20, max_duration_minutes=45, jezik=jezik
    ):
        audio = AudioSegment.from_file(input_path, format="mp3")
        max_file_size_bytes = max_file_size_mb * 1024 * 1024
        bitrate_kbps = 128
        max_duration_seconds_file_size = (max_file_size_bytes * 8) / (bitrate_kbps * 1000)
        max_duration_seconds_time = max_duration_minutes * 60
        max_duration_seconds = min(max_duration_seconds_file_size, max_duration_seconds_time)

        parts = []
        for i in range(0, len(audio), int(max_duration_seconds * 1000)):
            parts.append(audio[i : i + int(max_duration_seconds * 1000)])

        all_transcripts = []
        for idx, part in enumerate(parts):
            part_path = os.path.join(
                output_directory,
                f"{os.path.splitext(os.path.basename(input_path))[0]}_part{idx + 1}.mp3",
            )
            part.export(part_path, format="mp3", bitrate="128k")
            st.caption(f"Kreiram transkript {part_path}")
            all_transcripts.append(transcribe_audio(part_path, jezik))
        return " ".join(all_transcripts)

    def chunk_transcript(transkript_text, token_limit):
        words = transkript_text.split()
        chunks = []
        current_chunk = ""
        for word in words:
            if len((current_chunk + " " + word).split()) > token_limit:
                chunks.append(current_chunk.strip())
                current_chunk = word
            else:
                current_chunk += " " + word
        chunks.append(current_chunk.strip())
        return chunks

    convert_to_mp3(audio_file, "output.mp3")
    transcript = split_mp3_file("output.mp3", ".", jezik=jezik)

    st.caption("delim u delove po 1000 reci")
    chunks = chunk_transcript(transcript, 1000)
    st.caption(f"Broj delova je: {len(chunks)}")
    corrected_transcript = ""

    for i, chunk in enumerate(chunks):
        st.caption(f"Obradjujem {i + 1}. deo...")
        response = client.chat.completions.create(
            model=work_vars["names"]["openai_model"],
            temperature=0,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": chunk},
            ],
        )
        corrected_transcript += " " + response.choices[0].message.content.strip()

    return corrected_transcript


def delete_mp3_files(directory: str) -> None:
    """Delete leftover ``*.mp3`` files produced during transcription."""
    for mp3_file in glob.glob(os.path.join(directory, "*.mp3")):
        try:
            os.remove(mp3_file)
        except Exception as exc:  # noqa: BLE001
            st.info(f"Error deleting {mp3_file}: {exc}")
